from datetime import datetime, timedelta, timezone

import jwt
import pytest
from docx import Document as WordDocument
from fastapi import FastAPI
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.api.chat_routes import router as chat_router
from app.core.config import settings
from app.core.database import get_db
from app.ingestion.chunking import chunk_pages
from app.ingestion.parsers import extract_document_pages
from app.models import Base, Employee, Permission, Role
from app.models.audit_event import AuditEvent
from app.models.document import Document
from app.models.document_chunk import DocumentChunk
from app.models.document_processing_job import DocumentProcessingJob
from app.services import document_processing_service, document_rag_service
from app.services.document_faiss_service import DocumentFaissIndex
from app.services.document_access_service import can_access_document
from app.services.document_retrieval_service import retrieve_document_chunks
from app.services.embedding_service import (
    EmbeddingServiceError,
    embed_texts,
)
from app.services.document_job_queue import (
    claim_next_document_job,
    enqueue_document_processing,
    extend_document_job_lease,
    mark_document_job_failed,
)
from app.services.document_duplicate_service import (
    calculate_content_hash,
    calculate_text_fingerprint,
    find_duplicate_review_match,
    normalize_document_text,
    simhash_similarity,
)
from app.workers.document_worker import ensure_document_job_schema


@pytest.fixture
def pipeline_db(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "document_storage_dir", tmp_path)
    monkeypatch.setattr(
        settings,
        "document_faiss_index_path",
        tmp_path / "indexes" / "documents.faiss",
    )
    monkeypatch.setattr(
        settings,
        "jwt_secret_key",
        "document-pipeline-tests-secret-key",
    )

    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    test_session = sessionmaker(bind=engine, expire_on_commit=False)

    with test_session() as db:
        yield db, test_session, tmp_path

    Base.metadata.drop_all(engine)
    engine.dispose()


def _employee(
    db,
    employee_id: str,
    department: str = "Engineering",
    permissions: list[Permission] | None = None,
) -> Employee:
    role = Role(
        name=f"role-{employee_id}",
        permissions=permissions or [],
    )
    employee = Employee(
        employee_id=employee_id,
        full_name=employee_id,
        email=f"{employee_id.lower()}@example.test",
        department=department,
        role=role,
    )
    db.add(employee)
    db.flush()
    return employee


def _document(
    db,
    employee: Employee,
    name: str,
    access_scope: str,
    department: str | None,
    content: str,
    embedding: list[float],
) -> tuple[Document, DocumentChunk]:
    document = Document(
        original_filename=name,
        storage_key=f"{name}.txt",
        file_extension="txt",
        mime_type="text/plain",
        file_size=10,
        sha256=name.encode().hex()[:64].ljust(64, "0"),
        uploaded_by_employee_id=employee.employee_id,
        department=department,
        access_scope=access_scope,
        processing_status="ready",
    )
    db.add(document)
    db.flush()
    chunk = DocumentChunk(
        document_id=document.id,
        chunk_index=0,
        content=content,
        page_number=None,
        content_hash=content.encode().hex()[:64].ljust(64, "0"),
        embedding=embedding,
        embedding_model=settings.embedding_model,
        embedding_dimension=len(embedding),
    )
    db.add(chunk)
    db.flush()
    return document, chunk


def test_extract_docx_text_and_citation_metadata(tmp_path):
    file_path = tmp_path / "guide.docx"
    document = WordDocument()
    document.add_paragraph("First paragraph")
    document.add_table(rows=1, cols=2).rows[0].cells[0].text = "Table text"
    document.save(file_path)

    pages = extract_document_pages(file_path, "docx")

    assert pages == [
        {
            "source": "guide.docx",
            "page_number": None,
            "text": "First paragraph\n\nTable text | ",
        }
    ]


def test_txt_extraction_and_configurable_chunk_metadata(tmp_path):
    source = tmp_path / "guide.txt"
    source.write_text(
        " First paragraph.  \n\nSecond paragraph with more words. ",
        encoding="utf-8",
    )
    pages = extract_document_pages(source, "txt")
    chunks = chunk_pages(pages, chunk_size=25, overlap=5)

    assert pages[0]["page_number"] is None
    assert chunks
    assert [chunk["chunk_index"] for chunk in chunks] == list(range(len(chunks)))
    assert all(chunk["text"].strip() for chunk in chunks)
    assert all(chunk["source"] == "guide.txt" for chunk in chunks)
    assert all(chunk["page_number"] is None for chunk in chunks)
    assert chunk_pages(pages, chunk_size=25, overlap=5) == chunks


def test_pdf_extraction_preserves_page_numbers(tmp_path):
    source = tmp_path / "blank-pages.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    writer.add_blank_page(width=72, height=72)
    with source.open("wb") as file:
        writer.write(file)

    pages = extract_document_pages(source, "pdf")

    assert [page["page_number"] for page in pages] == [1, 2]
    assert all(not page["text"] for page in pages)
    assert chunk_pages(pages) == []


def test_corrupt_pdf_and_empty_document_are_not_processed(tmp_path):
    corrupt_pdf = tmp_path / "corrupt.pdf"
    corrupt_pdf.write_bytes(b"not a pdf")
    with pytest.raises((ValueError, RuntimeError)):
        extract_document_pages(corrupt_pdf, "pdf")

    empty_txt = tmp_path / "empty.txt"
    empty_txt.write_text(" \n\t ", encoding="utf-8")
    pages = extract_document_pages(empty_txt, "txt")
    assert chunk_pages(pages) == []


def test_embed_texts_validates_count_and_dimensions(monkeypatch):
    calls = []

    def fake_embed(*, model, input):
        calls.append((model, input))
        return {"embeddings": [[0.5, 0.5], [0.25, 0.75]]}

    monkeypatch.setattr("app.services.embedding_service._client.embed", fake_embed)

    result = embed_texts(["first", "second"])

    assert len(result) == 2
    assert calls == [(settings.embedding_model, ["first", "second"])]

    monkeypatch.setattr(
        "app.services.embedding_service._client.embed",
        lambda **_: {"embeddings": [[1.0]]},
    )
    with pytest.raises(ValueError, match="different number"):
        embed_texts(["first", "second"])

    def unavailable(**_):
        raise ConnectionError("offline")

    monkeypatch.setattr(
        "app.services.embedding_service._client.embed",
        unavailable,
    )
    with pytest.raises(EmbeddingServiceError, match="configured Ollama"):
        embed_texts(["first"])


def test_faiss_index_persists_ids_and_checks_query_dimensions(tmp_path):
    index = DocumentFaissIndex(tmp_path / "documents.faiss")
    index.rebuild([17, 42], [[1.0, 0.0], [0.0, 1.0]])

    results = DocumentFaissIndex(tmp_path / "documents.faiss").search(
        [0.9, 0.1]
    )

    assert [chunk_id for chunk_id, _ in results] == [17, 42]
    index.add(
        [17],
        [[1.0, 0.0]],
        model_name=settings.embedding_model,
    )
    assert DocumentFaissIndex(tmp_path / "documents.faiss").search(
        [0.9, 0.1]
    ) == results
    assert index.metadata_path.is_file()
    with pytest.raises(ValueError, match="dimensions"):
        index.search([1.0])


def test_faiss_rejects_mismatched_embedding_model(tmp_path, monkeypatch):
    index = DocumentFaissIndex(tmp_path / "documents.faiss")
    index.rebuild([17], [[1.0, 0.0]], model_name="model-a")
    monkeypatch.setattr(settings, "embedding_model", "model-b")

    with pytest.raises(ValueError, match="different embedding model"):
        index.search([1.0, 0.0])


def test_processing_uses_fresh_session_and_persists_ready_chunks(
    pipeline_db,
    monkeypatch,
):
    db, test_session, storage = pipeline_db
    employee = _employee(db, "UPLOADER")
    file_path = storage / "policy.txt"
    file_path.write_text("A short policy with useful information.", encoding="utf-8")
    document = Document(
        original_filename="policy.txt",
        storage_key="policy.txt",
        file_extension="txt",
        mime_type="text/plain",
        file_size=file_path.stat().st_size,
        sha256="a" * 64,
        uploaded_by_employee_id=employee.employee_id,
        access_scope="private",
        processing_status="uploaded",
    )
    db.add(document)
    db.commit()
    document_id = document.id

    opened_sessions = []

    def new_session():
        opened_sessions.append(True)
        return test_session()

    monkeypatch.setattr(document_processing_service, "SessionLocal", new_session)
    monkeypatch.setattr(
        "app.services.embedding_service._client.embed",
        lambda **_: {"embeddings": [[1.0, 0.0]]},
    )

    document_processing_service.process_document(document_id)

    db.expire_all()
    processed = db.get(Document, document_id)
    chunks = db.scalars(
        select(DocumentChunk).where(DocumentChunk.document_id == document_id)
    ).all()
    assert opened_sessions == [True]
    assert processed.processing_status == "ready"
    assert processed.content_hash == calculate_content_hash(
        "A short policy with useful information."
    )
    assert processed.text_fingerprint is None
    assert processed.duplicate_review_status == "none"
    assert len(chunks) == 1
    assert chunks[0].embedding == [1.0, 0.0]
    assert DocumentFaissIndex().index_path.is_file()
    assert DocumentFaissIndex().search([1.0, 0.0])[0][0] == chunks[0].id
    event = db.scalar(
        select(AuditEvent).where(
            AuditEvent.action == "document_processing_completed",
            AuditEvent.resource_id == str(document_id),
        )
    )
    assert event is not None
    assert event.details["version_number"] == 1


def test_successful_new_version_activates_only_after_indexing(
    pipeline_db,
    monkeypatch,
):
    db, test_session, storage = pipeline_db
    employee = _employee(db, "VERSION-ACTIVATION")
    previous, _ = _document(
        db,
        employee,
        "policy-v1",
        "private",
        None,
        "Previous ready policy text",
        [0.0, 1.0],
    )
    source = storage / "policy-v2.txt"
    source.write_text("New version policy text.", encoding="utf-8")
    version = Document(
        original_filename=source.name,
        storage_key=source.name,
        file_extension="txt",
        mime_type="text/plain",
        file_size=source.stat().st_size,
        sha256="f" * 64,
        document_series_id=previous.document_series_id,
        version_number=2,
        is_current=False,
        uploaded_by_employee_id=employee.employee_id,
        access_scope="private",
        processing_status="uploaded",
    )
    db.add(version)
    db.commit()
    version_id = version.id
    job = DocumentProcessingJob(
        document_id=version_id,
        status="running",
        attempts=2,
        lease_token="version-retry-token",
    )
    db.add(job)
    db.commit()
    job_id = job.id
    monkeypatch.setattr(document_processing_service, "SessionLocal", test_session)
    monkeypatch.setattr(
        "app.services.embedding_service._client.embed",
        lambda **_: {"embeddings": [[1.0, 0.0]]},
    )

    document_processing_service.process_document(
        version_id,
        job_id=job_id,
        lease_token="version-retry-token",
    )

    db.expire_all()
    previous = db.get(Document, previous.id)
    version = db.get(Document, version_id)
    assert previous.is_current is False
    assert previous.processing_status == "ready"
    assert version.is_current is True
    assert version.processing_status == "ready"
    assert version.processed_at is not None
    assert db.query(DocumentChunk).filter_by(document_id=version_id).count() == 1
    assert DocumentFaissIndex().index_path.is_file()
    activation = db.scalar(
        select(AuditEvent).where(
            AuditEvent.action == "document_version_activated",
            AuditEvent.resource_id == str(version_id),
        )
    )
    assert activation is not None
    retry_event = db.scalar(
        select(AuditEvent).where(
            AuditEvent.action == "document_processing_retried",
            AuditEvent.resource_id == str(version_id),
        )
    )
    assert retry_event is not None

    previous_source = storage / previous.storage_key
    previous_source.write_text(
        "Previous ready policy text.",
        encoding="utf-8",
    )
    document_processing_service.process_document(previous.id)
    db.expire_all()
    assert db.get(Document, previous.id).is_current is False
    assert db.get(Document, previous.id).processing_status == "ready"
    assert db.get(Document, version_id).is_current is True
    assert DocumentFaissIndex().index_path.is_file()
    assert len(DocumentFaissIndex().search([1.0, 0.0])) == 1

    document_processing_service.process_document(version_id)
    db.expire_all()
    assert db.query(DocumentChunk).filter_by(document_id=version_id).count() == 1


def test_failed_new_version_keeps_previous_ready_version_active(
    pipeline_db,
    monkeypatch,
):
    db, test_session, storage = pipeline_db
    employee = _employee(db, "VERSION-FAILURE")
    previous, _ = _document(
        db,
        employee,
        "stable-policy",
        "private",
        None,
        "Previous ready policy text",
        [0.0, 1.0],
    )
    source = storage / "empty-version.txt"
    source.write_text(" \n ", encoding="utf-8")
    version = Document(
        original_filename=source.name,
        storage_key=source.name,
        file_extension="txt",
        mime_type="text/plain",
        file_size=source.stat().st_size,
        sha256="e" * 64,
        document_series_id=previous.document_series_id,
        version_number=2,
        is_current=False,
        uploaded_by_employee_id=employee.employee_id,
        access_scope="private",
        processing_status="uploaded",
    )
    db.add(version)
    db.commit()
    version_id = version.id
    monkeypatch.setattr(document_processing_service, "SessionLocal", test_session)

    with pytest.raises(ValueError, match="no extractable text"):
        document_processing_service.process_document(version_id)

    db.expire_all()
    assert db.get(Document, previous.id).is_current is True
    assert db.get(Document, previous.id).processing_status == "ready"
    failed_version = db.get(Document, version_id)
    assert failed_version.is_current is False
    assert failed_version.processing_status == "failed"
    assert failed_version.processing_error_code == "document_processing_invalid"


def test_normalized_content_hash_ignores_case_and_whitespace():
    first = "Employee Leave Policy\nAnnual leave requires approval."
    second = "  EMPLOYEE   LEAVE POLICY annual leave requires approval. "

    assert normalize_document_text(first) == normalize_document_text(second)
    assert calculate_content_hash(first) == calculate_content_hash(second)


def test_text_fingerprint_flags_near_duplicate_for_review(pipeline_db, monkeypatch):
    db, _, _ = pipeline_db
    monkeypatch.setattr(settings, "document_near_duplicate_threshold", 0.98)
    employee = _employee(db, "DUPLICATE-REVIEW")
    original_text = (
        "The company leave policy provides annual leave employees submit "
        "requests through the portal with manager approval and human resources "
        "tracks balances. "
    ) * 4
    changed_text = original_text.replace("annual leave", "annual vacation", 1)
    first = Document(
        original_filename="leave-policy-2025.pdf",
        storage_key="first.pdf",
        file_extension="pdf",
        mime_type="application/pdf",
        file_size=20,
        sha256="1" * 64,
        uploaded_by_employee_id=employee.employee_id,
        access_scope="private",
        processing_status="ready",
    )
    second = Document(
        original_filename="leave-policy-2026.pdf",
        storage_key="second.pdf",
        file_extension="pdf",
        mime_type="application/pdf",
        file_size=20,
        sha256="2" * 64,
        uploaded_by_employee_id=employee.employee_id,
        access_scope="private",
        processing_status="processing",
    )
    db.add_all([first, second])
    db.flush()

    find_duplicate_review_match(db, first, original_text)
    db.commit()
    find_duplicate_review_match(db, second, changed_text)
    db.commit()

    assert calculate_text_fingerprint(original_text) is not None
    assert simhash_similarity(
        calculate_text_fingerprint(original_text),
        calculate_text_fingerprint(changed_text),
    ) >= 0.98
    assert second.duplicate_review_status == "near_duplicate"
    assert second.duplicate_similarity >= 0.98


def test_exact_extracted_content_is_flagged_without_reusing_record(pipeline_db):
    db, _, _ = pipeline_db
    first_employee = _employee(db, "CONTENT-DUPLICATE-A")
    second_employee = _employee(db, "CONTENT-DUPLICATE-B", department="HR")
    first = Document(
        original_filename="employee_policy.pdf",
        storage_key="original.pdf",
        file_extension="pdf",
        mime_type="application/pdf",
        file_size=20,
        sha256="3" * 64,
        uploaded_by_employee_id=first_employee.employee_id,
        access_scope="private",
        processing_status="ready",
    )
    second = Document(
        original_filename="employee_policy_final.pdf",
        storage_key="final.pdf",
        file_extension="pdf",
        mime_type="application/pdf",
        file_size=21,
        sha256="4" * 64,
        uploaded_by_employee_id=second_employee.employee_id,
        access_scope="private",
        processing_status="processing",
    )
    text = "The employee policy contains enough substantive text for a document."
    db.add_all([first, second])
    db.flush()

    find_duplicate_review_match(db, first, text)
    db.commit()
    first_id = first.id
    find_duplicate_review_match(db, second, text.upper())
    db.commit()

    assert first.id == first_id
    assert second.id != first.id
    assert second.content_hash == first.content_hash
    assert second.duplicate_review_status == "exact_content"
    assert second.duplicate_similarity == 1.0
    assert db.query(Document).count() == 2
    assert can_access_document(first, first_employee)
    assert not can_access_document(first, second_employee)
    assert can_access_document(second, second_employee)
    assert not can_access_document(second, first_employee)


def test_short_text_does_not_produce_near_duplicate_fingerprint():
    assert calculate_text_fingerprint("brief note") is None


def test_processing_failure_marks_document_failed(
    pipeline_db,
    monkeypatch,
    caplog,
):
    db, test_session, storage = pipeline_db
    employee = _employee(db, "BROKEN-UPLOADER")
    file_path = storage / "empty.txt"
    file_path.write_text("   \n", encoding="utf-8")
    document = Document(
        original_filename="empty.txt",
        storage_key="empty.txt",
        file_extension="txt",
        mime_type="text/plain",
        file_size=file_path.stat().st_size,
        sha256="b" * 64,
        uploaded_by_employee_id=employee.employee_id,
        access_scope="private",
        processing_status="uploaded",
    )
    db.add(document)
    db.commit()
    document_id = document.id
    monkeypatch.setattr(
        document_processing_service,
        "SessionLocal",
        test_session,
    )

    with pytest.raises(ValueError, match="no extractable text"):
        document_processing_service.process_document(document_id)

    db.expire_all()
    assert db.get(Document, document_id).processing_status == "failed"
    assert "Document processing failed" in caplog.text
    event = db.scalar(
        select(AuditEvent).where(
            AuditEvent.action == "document_processing_failed",
            AuditEvent.resource_id == str(document_id),
        )
    )
    assert event is not None
    assert event.outcome == "failure"
    assert event.details["error_code"] == "document_processing_invalid"


def test_job_queue_retries_with_backoff_and_recovers_expired_lease(
    pipeline_db,
    monkeypatch,
):
    db, _, _ = pipeline_db
    monkeypatch.setattr(settings, "document_job_max_attempts", 3)
    monkeypatch.setattr(settings, "document_job_retry_base_seconds", 5)
    monkeypatch.setattr(settings, "document_job_lease_seconds", 10)
    employee = _employee(db, "QUEUE-UPLOADER")
    document = Document(
        original_filename="queued.txt",
        storage_key="queued.txt",
        file_extension="txt",
        mime_type="text/plain",
        file_size=5,
        sha256="c" * 64,
        uploaded_by_employee_id=employee.employee_id,
        access_scope="private",
        processing_status="uploaded",
    )
    db.add(document)
    db.flush()
    job = enqueue_document_processing(db, document.id)
    db.commit()
    job_id = job.id
    document_id = document.id
    now = datetime.now(timezone.utc) + timedelta(seconds=1)

    first_claim = claim_next_document_job(db, now=now)
    assert first_claim is not None
    assert first_claim.attempts == 1
    mark_document_job_failed(
        db,
        job_id,
        first_claim.attempts,
        first_claim.lease_token,
        RuntimeError("temporary model outage"),
        now=now,
    )
    db.expire_all()
    job = db.get(DocumentProcessingJob, job_id)
    assert job.status == "queued"
    assert job.available_at.replace(tzinfo=timezone.utc) == (
        now + timedelta(seconds=5)
    )

    assert claim_next_document_job(
        db,
        now=now + timedelta(seconds=4),
    ) is None
    second_claim = claim_next_document_job(
        db,
        now=now + timedelta(seconds=5),
    )
    assert second_claim is not None
    assert second_claim.attempts == 2
    mark_document_job_failed(
        db,
        job_id,
        first_claim.attempts,
        first_claim.lease_token,
        RuntimeError("stale worker"),
        now=now + timedelta(seconds=5),
    )
    db.expire_all()
    assert db.get(DocumentProcessingJob, job_id).status == "running"

    third_claim = claim_next_document_job(
        db,
        now=now + timedelta(seconds=16),
    )
    assert third_claim is not None
    assert third_claim.attempts == 3
    mark_document_job_failed(
        db,
        job_id,
        third_claim.attempts,
        third_claim.lease_token,
        RuntimeError("permanent failure"),
        now=now + timedelta(seconds=16),
    )
    db.expire_all()

    assert db.get(DocumentProcessingJob, job_id).status == "failed"
    assert db.get(Document, document_id).processing_status == "failed"


def test_job_lease_heartbeat_extends_only_current_lease(pipeline_db):
    db, _, _ = pipeline_db
    employee = _employee(db, "HEARTBEAT-UPLOADER")
    document = Document(
        original_filename="heartbeat.txt",
        storage_key="heartbeat.txt",
        file_extension="txt",
        mime_type="text/plain",
        file_size=5,
        sha256="e" * 64,
        uploaded_by_employee_id=employee.employee_id,
        access_scope="private",
        processing_status="uploaded",
    )
    db.add(document)
    db.flush()
    job = enqueue_document_processing(db, document.id)
    db.commit()
    claimed = claim_next_document_job(db)
    assert claimed is not None

    assert extend_document_job_lease(
        db,
        claimed.job_id,
        claimed.lease_token,
    )
    assert not extend_document_job_lease(
        db,
        claimed.job_id,
        "stale-lease-token",
    )


def test_worker_claims_and_completes_persisted_job(
    pipeline_db,
    monkeypatch,
):
    db, test_session, _ = pipeline_db
    employee = _employee(db, "WORKER-UPLOADER")
    document = Document(
        original_filename="worker.txt",
        storage_key="worker.txt",
        file_extension="txt",
        mime_type="text/plain",
        file_size=5,
        sha256="f" * 64,
        uploaded_by_employee_id=employee.employee_id,
        access_scope="private",
        processing_status="uploaded",
    )
    db.add(document)
    db.flush()
    job = enqueue_document_processing(db, document.id)
    db.commit()
    job_id = job.id
    monkeypatch.setattr(
        "app.core.database.SessionLocal",
        test_session,
    )
    monkeypatch.setattr(
        document_processing_service,
        "process_document",
        lambda document_id, **kwargs: None,
    )

    from app.services.document_job_queue import process_one_document_job

    assert process_one_document_job()

    db.expire_all()
    completed = db.get(DocumentProcessingJob, job_id)
    assert completed.status == "completed"
    assert completed.attempts == 1
    assert completed.lease_token is None


def test_requeue_refuses_a_job_with_an_active_lease(pipeline_db):
    db, _, _ = pipeline_db
    employee = _employee(db, "ACTIVE-UPLOADER")
    document = Document(
        original_filename="active.txt",
        storage_key="active.txt",
        file_extension="txt",
        mime_type="text/plain",
        file_size=5,
        sha256="d" * 64,
        uploaded_by_employee_id=employee.employee_id,
        access_scope="private",
        processing_status="uploaded",
    )
    db.add(document)
    db.flush()
    enqueue_document_processing(db, document.id)
    db.commit()
    claimed = claim_next_document_job(db)
    assert claimed is not None

    from app.services.document_job_queue import JobAlreadyRunningError

    with pytest.raises(JobAlreadyRunningError):
        enqueue_document_processing(db, document.id)


def test_worker_fails_fast_when_queue_migration_is_missing():
    engine = create_engine("sqlite://")

    with pytest.raises(
        RuntimeError,
        match=r"Run `alembic upgrade head`",
    ):
        ensure_document_job_schema(engine)

    engine.dispose()


def test_retrieval_filters_inaccessible_document_text_before_rag(
    pipeline_db,
    monkeypatch,
):
    db, _, _ = pipeline_db
    reader = _employee(db, "READER")
    owner = _employee(db, "OTHER")
    _, company_chunk = _document(
        db, owner, "company", "company", None, "company context", [0.9, 0.1]
    )
    _, department_chunk = _document(
        db,
        owner,
        "engineering",
        "department",
        "Engineering",
        "department context",
        [0.8, 0.2],
    )
    _, private_chunk = _document(
        db,
        owner,
        "secret",
        "private",
        None,
        "private secret context",
        [1.0, 0.0],
    )
    _, own_chunk = _document(
        db,
        reader,
        "own",
        "private",
        None,
        "own private context",
        [0.7, 0.3],
    )
    db.commit()
    DocumentFaissIndex().rebuild(
        [
            company_chunk.id,
            department_chunk.id,
            private_chunk.id,
            own_chunk.id,
        ],
        [[0.9, 0.1], [0.8, 0.2], [1.0, 0.0], [0.7, 0.3]],
    )
    monkeypatch.setattr(
        "app.services.embedding_service._client.embed",
        lambda **_: {"embeddings": [[1.0, 0.0]]},
    )

    retrieved = retrieve_document_chunks(db, reader, "company policy")
    model_messages = {}

    def fake_chat(*, model, messages):
        model_messages["model"] = model
        model_messages["messages"] = messages
        return {"message": {"content": "Grounded answer [company.txt]."}}

    monkeypatch.setattr(document_rag_service.ollama, "chat", fake_chat)
    answer = document_rag_service.answer_document_question(
        db,
        reader,
        "What is in the documents?",
    )
    prompt = model_messages["messages"][1]["content"]

    assert {chunk["content"] for chunk in retrieved} == {
        "company context",
        "department context",
        "own private context",
    }
    assert "private secret context" not in prompt
    assert len(answer["sources"]) == 3
    assert answer["answer"].startswith("Grounded answer")


def test_retrieval_excludes_chunks_from_historical_versions(
    pipeline_db,
    monkeypatch,
):
    db, _, _ = pipeline_db
    employee = _employee(db, "VERSION-READER")
    old_document, old_chunk = _document(
        db,
        employee,
        "policy-v1",
        "private",
        None,
        "outdated policy text",
        [1.0, 0.0],
    )
    old_document.is_current = False
    db.flush()
    current_document, current_chunk = _document(
        db,
        employee,
        "policy-v2",
        "private",
        None,
        "current policy text",
        [0.9, 0.1],
    )
    current_document.document_series_id = old_document.document_series_id
    current_document.version_number = 2
    db.commit()
    DocumentFaissIndex().rebuild(
        [old_chunk.id, current_chunk.id],
        [[1.0, 0.0], [0.9, 0.1]],
    )
    monkeypatch.setattr(
        "app.services.embedding_service._client.embed",
        lambda **_: {"embeddings": [[1.0, 0.0]]},
    )

    retrieved = retrieve_document_chunks(db, employee, "policy")

    assert [chunk["document_id"] for chunk in retrieved] == [current_document.id]
    assert "outdated policy text" not in {
        chunk["content"] for chunk in retrieved
    }


def test_chat_endpoint_requires_documents_read(pipeline_db, monkeypatch):
    db, _, _ = pipeline_db
    read_permission = Permission(code="documents:read", description="Read")
    employee = _employee(db, "CHATTER")
    db.commit()
    monkeypatch.setattr(
        "app.api.chat_routes.answer_document_question",
        lambda *_: {"answer": "Answer", "sources": []},
    )

    app = FastAPI()
    app.include_router(chat_router)
    def override_db():
        yield db

    app.dependency_overrides[get_db] = override_db
    token = jwt.encode(
        {
            "sub": employee.employee_id,
            "exp": datetime.now(timezone.utc) + timedelta(minutes=5),
        },
        settings.jwt_secret_key,
        algorithm=settings.jwt_algorithm,
    )
    with TestClient(app) as client:
        headers = {"Authorization": f"Bearer {token}"}
        denied_response = client.post(
            "/chat/ask",
            headers=headers,
            json={"message": "Question"},
        )
        employee.role.permissions.append(read_permission)
        db.commit()
        response = client.post(
            "/chat/ask",
            headers=headers,
            json={"message": "Question"},
        )

    assert denied_response.status_code == 403
    assert response.status_code == 200
    assert response.json()["answer"] == "Answer"
