"""Document parsers."""

from pathlib import Path
from typing import TypedDict

from docx import Document as WordDocument
from pypdf import PdfReader


class ExtractedPage(TypedDict):
    source: str
    page_number: int | None
    text: str


def extract_pdf_pages(pdf_path: str | Path) -> list[ExtractedPage]:
    """
    Extract text from a PDF, preserving its filename and page numbers.

    Page numbers are 1-based for human-readable citations.
    Pages without extractable text are retained with empty text.
    """
    path = Path(pdf_path)

    if not path.is_file():
        raise FileNotFoundError(f"PDF file not found: {path}")

    try:
        reader = PdfReader(str(path))

        if reader.is_encrypted:
            raise ValueError(
                f"Encrypted PDFs are not supported: {path.name}"
            )

        extracted_pages: list[ExtractedPage] = []

        for page_number, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""

            # Normalize line endings and remove null characters.
            text = text.replace("\x00", "").replace("\r\n", "\n")
            text = text.replace("\r", "\n").strip()

            extracted_pages.append(
                {
                    "source": path.name,
                    "page_number": page_number,
                    "text": text,
                }
            )

        return extracted_pages

    except ValueError:
        raise
    except Exception as exc:
        raise RuntimeError(
            f"Failed to extract text from {path.name}: {exc}"
        ) from exc


def extract_document_pages(
    file_path: str | Path,
    file_extension: str,
) -> list[ExtractedPage]:
    """Extract text from supported document formats with citation metadata."""
    path = Path(file_path)
    if not path.is_file():
        raise FileNotFoundError(f"Document file not found: {path}")

    extension = file_extension.lower().lstrip(".")
    if extension == "pdf":
        return extract_pdf_pages(path)

    if extension == "docx":
        try:
            document = WordDocument(path)
            paragraphs = [
                paragraph.text
                for paragraph in document.paragraphs
                if paragraph.text.strip()
            ]
            for table in document.tables:
                paragraphs.extend(
                    " | ".join(cell.text for cell in row.cells)
                    for row in table.rows
                )
        except Exception as exc:
            raise RuntimeError(
                f"Failed to extract text from {path.name}: {exc}"
            ) from exc

        text = "\n\n".join(paragraphs)
        return [{"source": path.name, "page_number": None, "text": text}]

    if extension == "txt":
        try:
            text = path.read_text(encoding="utf-8")
        except (OSError, UnicodeError) as exc:
            raise RuntimeError(
                f"Failed to read text from {path.name}: {exc}"
            ) from exc
        return [{"source": path.name, "page_number": None, "text": text}]

    raise ValueError(f"Unsupported document extension: {file_extension}")
